# -*- coding: utf-8 -*-
"""生成两份 Word 交付物（python-docx，零模板依赖）：
  docs/industry-analysis.docx        行业分析
  docs/design-proposal-quote.docx   设计方案 + 报价单
运行：python scripts/gen_docs.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE = os.path.dirname(__file__)
OUT = os.path.abspath(os.path.join(HERE, "..", "docs"))
NAVY = RGBColor(0x1E, 0x3A, 0x8A)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GREY = RGBColor(0x6B, 0x72, 0x80)

def new_doc(title):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(11)
    # 标题
    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
    sub = d.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run("北辰在线学院 · AI 学习支持与辅导协作  |  交付物  |  2026-10-29")
    rs.font.size = Pt(10); rs.font.color.rgb = GREY
    d.add_paragraph()
    return d

def h1(d, text):
    p = d.add_heading(level=1)
    r = p.add_run(text); r.font.color.rgb = NAVY; r.bold = True; r.font.size = Pt(15)
    return p

def h2(d, text):
    p = d.add_heading(level=2)
    r = p.add_run(text); r.font.color.rgb = BLUE; r.bold = True; r.font.size = Pt(13)
    return p

def para(d, text, bold=False, size=11, color=None):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p

def bullets(d, items):
    for it in items:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(it)

def table(d, headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True; run.font.size = Pt(10.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(10)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t

# ========================= 行业分析 =========================
def build_industry():
    d = new_doc("行业分析：AI 学习支持与辅导协作")
    h1(d, "1. 摘要")
    para(d, "北辰在线学院拟在“学习支持与辅导协作”环节引入 AI，目标把支持请求首次响应从 26 小时压缩到 8 小时、重复求助率从 35% 降到 15%，并在两个班级试点中将求助漏单控制在 2% 以内。本分析从在线教育行业、AI 教育应用、学生支持共性痛点、可比技术实践、数据与接口条件、切入点合理性、风险结论七个维度，论证该方向在行业内的成熟度与本项目落地的特殊约束。")

    h1(d, "2. 在线教育行业概况")
    h2(d, "2.1 市场规模与增长")
    para(d, "中国在线教育自 2020 年后进入存量精细化运营阶段，职业教育与高等教育继续教育保持稳健增长。机构竞争焦点从“获客”转向“完课率与学习效果”，这意味着学习支持（答疑、辅导、预警）正成为差异化能力，而非边缘职能。")
    h2(d, "2.2 政策与监管环境")
    bullets(d, [
        "“双减”后学科类培训受限，但高等教育、职业教育与成人继续教育不在限制范围，北辰所在的高等继续教育赛道政策空间充足。",
        "《个人信息保护法》《数据安全法》要求学习数据“最小必要、可追溯”，直接对应本项目中“建议必须附行为证据、审计可重建处理链”的硬约束。",
        "教育评价改革强调“不唯分数”，与本项目“不自动给学生定性或评分”的红线一致。",
    ])

    h1(d, "3. AI 在教育领域的应用现状")
    h2(d, "3.1 学习分析与学情预警")
    para(d, "以 Canvas、Blackboard 为代表的学习管理系统已内置学习分析仪表盘，通过登录率、视频时长、作业提交等信号做掉队预警。但行业共识是：活跃度 ≠ 理解度，单纯时长指标误报率高——这正是北辰报告中“视频播放时长高但挂着播放器”痛点的行业共性。")
    h2(d, "3.2 智能答疑与辅导")
    bullets(d, [
        "大模型出现后，智能答疑从“关键词检索”升级为“语义理解 + 多轮对话”，能处理作业评论里的自然语言求助。",
        "主流做法均保留“人工复核”：AI 生成建议，教师/助教确认后再触达学生，避免错误建议污染原记录。",
        "“不自动定性评分”已成为合规底线，海外与国内头部机构均将评分动作保留给人工。",
    ])
    h2(d, "3.3 教学运营与工单化支持")
    para(d, "学生支持正在从“群聊接龙”走向“工单化 SLA 管理”（参考 Zendesk、Intercom 在教育客服的落地）。工单化带来统一编号、责任交接、SLA 计时与漏单统计，正好解决北辰“入口分散、交接无随单说明、漏单无基线”的核心问题。")

    h1(d, "4. 学生支持 / 辅导协作的共性痛点")
    bullets(d, [
        "入口分散：求助散落群聊、作业评论、邮件，无统一受理渠道，连续求助易被漏接。",
        "交接断层：关键说明留在个人便签或群里，不随单据交接，接手人把连续求助当新问题。",
        "数据孤岛：LMS、直播、工单系统账号与状态各自独立，缺统一业务视图。",
        "口径不一：运营/讲师/助教各记一套指标，重复求助与交接等待易被重复计入。",
        "反馈失真：以通知发送量代替实际完成量，误判例子无沉淀，规则迭代无依据。",
    ])

    h1(d, "5. 可比技术实践")
    table(d, ["实践 / 产品", "做法", "对本项目的启示"], [
        ["Georgia Tech Jill Watson", "用 AI 助教回答 Piazza 论坛高频问题，人工兜底复杂诉求", "AI 先做高频识别，复杂/敏感转人工，与本设计一致"],
        ["Knewton / 自适应学习", "基于行为数据做掌握度诊断", "必须先用行为证据而非时长，呼应 REQ-3 数据校验"],
        ["学而思 / 作业帮 AI 答疑", "作业图片/文本识别 + 大模型讲解，人工复核", "多源识别 + 人工确认后触达的学生侧范式"],
        ["Canvas 学习分析", "统一事件流 + 预警仪表盘", "统一业务编号/版本/时间戳是前提"],
        ["Zendesk 教育客服", "工单化 SLA、交接确认、审计留痕", "漏单率/接手确认的运营机制可借鉴"],
    ], widths=[1.6, 2.6, 2.4])

    h1(d, "6. 数据与接口条件（本项目特殊性）")
    bullets(d, [
        "三套系统（LMS、直播平台、辅导工单）账号权限独立维护，业务状态尚无统一视图，集成需逐系统申请授权。",
        "作业补交/撤回缺统一编号，学习活跃度不等于理解，进模型前必须做主编号/事件时间/版本校验，歧义记录进人工修复队列。",
        "合规红线：不依据敏感属性判断能力、不自动定性评分、审计可重建处理链——这些不是本项目独有，而是行业通行约束。",
    ])

    h1(d, "7. 本项目 AI 切入点的行业合理性")
    para(d, "行业成熟实践已验证“多源识别 + 人工复核 + 工单化交接 + 审计留痕”路径可行。北辰的独特价值在于：把切入点收敛到“支持需求识别与交接”这一漏单主因（REQ-1/REQ-2），而非直接替换现有流程，符合“小范围并行试点、不直接替换”的稳妥路径，也最契合 9 万元预算与两班级试点的现实约束。")

    h1(d, "8. 风险与可行性结论")
    bullets(d, [
        "可行性：高。核心能力（识别、合并、建议、复核）均有成熟行业范式，MVP 可用规则引擎先行验证闭环。",
        "主要风险：数据口径未统一（漏单基线、合并规则、统计周期均待确认），需在试点前固化，否则验收不可信。",
        "结论：建议优先在“识别与合并”环节做最小可验证方案，以漏单率<2% 与人工确认后联系为硬验收线。",
    ])

    h1(d, "9. 参考来源")
    bullets(d, [
        "《北辰在线学院 · 业务需求发现报告》（内部，2026-10-29）",
        "fde-simulation 框架：workflow-decomposition / outcome-risk-matrix / 4-source-convergence",
        "行业参考：Jill Watson (Georgia Tech)、Knewton、Canvas Analytics、Zendesk for Education",
    ])
    path = os.path.join(OUT, "industry-analysis.docx")
    d.save(path)
    return path

# ========================= 设计方案 + 报价单 =========================
def build_design():
    d = new_doc("设计方案与报价单：AI 学习支持系统")
    h1(d, "1. 项目概述")
    para(d, "本项目为北辰在线学院“学习支持与辅导协作”环节建设一套 AI 辅助系统：在群聊/作业评论/LMS 等多源入口自动识别支持需求、按学生合并连续求助并生成建议，由助教人工复核确认后再联系学生，配套随单交接、人工兜底与审计链。首期在两个班级试点，预算 9 万元（含集成、培训与验证）。")

    h1(d, "2. 业务目标与验收")
    table(d, ["指标", "现状基线", "目标", "验收口径"], [
        ["首次响应时长", "26 小时", "8 小时", "从求助产生到人工确认联系"],
        ["重复求助率", "35%", "15%", "同一学生连续求助合并计一单"],
        ["求助漏单率", "无基线（待回放）", "< 2%", "两班级试点，人工确认后联系"],
    ], widths=[1.4, 1.8, 1.3, 2.5])

    h1(d, "3. 系统架构设计")
    h2(d, "3.1 总体架构（四层）")
    bullets(d, [
        "采集层：对接群聊/作业评论/LMS/直播，统一事件流，带业务编号 + 时间戳 + 版本。",
        "AI 层：需求信号识别、连续求助合并、建议生成（显示原记录 + 待确认字段）。",
        "人工兜底层：助教复核（确认/转人工/暂停），接手确认闭环。",
        "审计层：每次处理留存编号/版本/事实/建议/人工改动/签核人，可重建处理链。",
    ])
    h2(d, "3.2 AI 模块")
    para(d, "采用“规则引擎 + 可选大模型”双模：规则引擎（关键词/信号）保障离线可跑与可解释，大模型作为可插拔增强。建议必须附原记录与待确认字段，不自动定性或评分。")
    h2(d, "3.3 人工兜底与升级路径")
    bullets(d, [
        "触发转人工：缺依据、两版本冲突、关键系统不可用、命中红线（成绩/处分等）。",
        "转人工须标注触发原因、接手人、确认时间，不覆盖原记录。",
        "超过约定时限由课程运营员升级至学习服务主任。",
    ])
    h2(d, "3.4 数据流与审计链")
    para(d, "接口中断时沿用带时间戳的只读快照并提示“暂未更新”，重连后按业务编号去重补齐，不重复派单，现有人工流程继续运行。")

    h1(d, "4. 实施阶段与里程碑")
    table(d, ["阶段", "周期", "交付物", "里程碑"], [
        ["P0 需求与数据校验", "第 1–3 周", "漏单基线、合并口径、数据校验规则", "顾明岚书面确认判定规则"],
        ["P1 AI 引擎开发", "第 3–7 周", "识别/合并/建议模块 + MVP", "MVP 闭环演示"],
        ["P2 交接与兜底", "第 6–9 周", "随单交接、接手确认、审计链", "转人工/暂停路径走通"],
        ["P3 集成与试点", "第 9–12 周", "三系统只读集成、两班级试点", "漏单<2%、人工确认后联系"],
    ], widths=[1.7, 1.1, 2.4, 1.8])

    h1(d, "5. 报价单")
    table(d, ["#", "模块 / 工作项", "内容说明", "人日", "报价(元)"], [
        ["1", "需求与数据校验", "漏单基线回放、连续求助合并口径、主编号/事件时间/版本校验", "12", "18,000"],
        ["2", "AI 识别与合并引擎", "多源采集解析、需求信号识别、按学生合并工单、建议生成", "16", "24,000"],
        ["3", "随单交接与接手确认", "业务编号/版本/状态/未解决事项/接手人，接手确认闭环", "8", "12,000"],
        ["4", "人工兜底 / 审计链", "暂停/转人工/复核，触发原因+接手人+确认时间，审计重建", "8", "12,000"],
        ["5", "系统集成（只读）", "LMS/直播/工单授权、字段映射、小时级快照、中断降级", "10", "15,000"],
        ["6", "试点培训与验证", "两班级试点、培训、误判例子收集、规则版本迭代", "4", "6,000"],
        ["7", "项目管理与运维", "进度/质量/验收对接、3 个月运维与复盘", "2", "3,000"],
        ["—", "合计", "—", "60", "90,000"],
    ], widths=[0.4, 1.6, 3.0, 0.6, 1.0])
    para(d, "说明：报价含集成、培训与验证（对应预算 9 万元）。收益须扣除维护与人工复核成本，不把节省工时直接等同现金回款。", size=10, color=GREY)

    h1(d, "6. 验收标准与停止条件")
    bullets(d, [
        "验收硬线：两班级试点漏单 < 2%；每条建议均可追溯到行为证据；人工确认后再联系学生。",
        "停止条件（任一触发即暂停并复盘）：建议无行为证据、依据敏感属性判断能力、关键系统不可用、触及合规红线。",
    ])

    h1(d, "7. 假设与约束")
    bullets(d, [
        "预算 9 万元是否覆盖编号改造与集成尚未核实，报价按首期范围估算。",
        "漏单基线、合并规则、统计周期、复核时限等口径须由业务负责人书面确认后方可用于验收。",
        "试点原始资料留存 90 天为暂定，真实期限按组织规则另行确定。",
    ])
    path = os.path.join(OUT, "design-proposal-quote.docx")
    d.save(path)
    return path

if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    p1 = build_industry()
    p2 = build_design()
    print("written:", p1)
    print("written:", p2)
